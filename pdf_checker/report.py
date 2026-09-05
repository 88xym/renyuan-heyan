# -*- coding: utf-8 -*-
"""报告输出：JSON（结构化）+ Excel（逐字段明细）+ 控制台摘要。"""

from __future__ import annotations

import json
import os
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .validate import FileResult

_PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
_PASS_FONT = Font(color="006100")
_FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
_FAIL_FONT = Font(color="9C0006")
_SKIP_FILL = PatternFill("solid", fgColor="FFEB9C")
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_HEADER_FONT = Font(bold=True)


def to_dict(result: FileResult) -> Dict[str, Any]:
    return {
        "file": result.file_name,
        "overall": "通过" if result.overall else "不通过",
        "fields": [
            {
                "label": f.label,
                "value": f.value,
                "found": f.found,
                "required": f.required,
                "passed": f.passed,
                "checks": [
                    {"type": r.rule_type, "passed": r.passed, "message": r.message}
                    for r in f.rules
                ],
            }
            for f in result.fields
        ],
        "cross_checks": [
            {"name": c.name, "passed": c.passed, "message": c.message}
            for c in result.cross_checks
        ],
        "issues": result.issues,
    }


def write_json(results: List[FileResult], out_path: str) -> None:
    data = [to_dict(r) for r in results]
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def write_excel(results: List[FileResult], out_path: str) -> None:
    """输出两个 Sheet：汇总 + 明细。"""
    wb = Workbook()

    # ---- Sheet1 汇总 ----
    ws = wb.active
    ws.title = "汇总"
    headers = ["文件", "字段数", "通过字段", "失败字段", "整体结果", "问题清单"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in results:
        n_fields = len(r.fields)
        n_pass = sum(1 for f in r.fields if f.passed)
        n_fail = n_fields - n_pass
        issues = "；".join(r.issues) if r.issues else ""
        row = [r.file_name, n_fields, n_pass, n_fail, "通过" if r.overall else "不通过", issues]
        ws.append(row)
        last = ws.max_row
        ws.cell(row=last, column=5).fill = _PASS_FILL if r.overall else _FAIL_FILL
        ws.cell(row=last, column=5).font = _PASS_FONT if r.overall else _FAIL_FONT

    widths = [32, 8, 9, 9, 10, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---- Sheet2 明细 ----
    ws2 = wb.create_sheet("字段明细")
    headers2 = ["文件", "字段(标签)", "提取值", "必填", "检查项", "结果", "说明"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in results:
        for f in r.fields:
            checks = f.rules or [type("R", (), {"rule_type": "-", "passed": f.passed, "message": ""})()]
            for chk in checks:
                status = "通过" if chk.passed else "不通过"
                ws2.append(
                    [r.file_name, f.label, f.value or "", "是" if f.required else "否",
                     chk.rule_type, status, chk.message]
                )
                last = ws2.max_row
                ws2.cell(row=last, column=6).fill = _PASS_FILL if chk.passed else _FAIL_FILL
                ws2.cell(row=last, column=6).font = _PASS_FONT if chk.passed else _FAIL_FONT
        for c in r.cross_checks:
            status = "通过" if c.passed else "不通过"
            ws2.append([r.file_name, "（跨字段）", "", "-", c.name, status, c.message])
            last = ws2.max_row
            ws2.cell(row=last, column=6).fill = _PASS_FILL if c.passed else _FAIL_FILL
            ws2.cell(row=last, column=6).font = _PASS_FONT if c.passed else _FAIL_FONT

    widths2 = [32, 14, 32, 6, 12, 9, 50]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)


def print_console(results: List[FileResult]) -> None:
    print("\n" + "=" * 60)
    for r in results:
        mark = "通过" if r.overall else "不通过"
        print(f"[{mark}] {r.file_name}")
        if r.issues:
            for issue in r.issues:
                print(f"    - {issue}")
    print("=" * 60)
